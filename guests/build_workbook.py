"""Rebuild the wedding planning workbook: all six tabs, olive green, new guests merged in."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from datetime import date

OUT = "/home/user/Claude/guests/Wedding_Planning_Workbook.xlsx"

# ---------------------------------------------------------------- olive palette
DEEP   = "4B5320"   # tab title bars, header rows
MID    = "6B7A3A"   # section labels, totals, accents
BAND   = "E7EAD4"   # banded fill, odd
BAND2  = "F4F6EA"   # banded fill, even
LINE   = "BFC79E"   # borders
TEXT   = "2F3517"   # body text
CREAM  = "FBFCF6"   # page ground

thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY = '"$"#,##0'

def title_block(ws, title, subtitle, span):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    c = ws.cell(row=1, column=1, value=title)
    c.font = Font(bold=True, size=16, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=DEEP)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
    s = ws.cell(row=2, column=1, value=subtitle)
    s.font = Font(size=10, italic=True, color=MID)
    s.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 18
    for col in range(1, span + 1):
        ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor=DEEP)

def header_row(ws, row, headers):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=DEEP)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 30

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

def band(ws, row, ncols, shade):
    fill = PatternFill("solid", fgColor=BAND if shade else BAND2)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center")
        if not cell.font.bold:
            cell.font = Font(color=TEXT, size=10)

wb = Workbook()

# ================================================================ GUEST LIST
gl = wb.active
gl.title = "Guest List"
GL_HEAD = ["#", "Guest / Party", "Side", "Category", "Household / Group", "Party Size",
           "+1?", "Mailing Address", "Phone", "Email", "Save-the-Date Sent", "Invite Sent",
           "RSVP", "# Attending", "Meal Choice", "Gift Received", "Thank-You Sent", "Notes"]

T1 = [("Mom & Dad","Holliday - Immediate",2,"No","Parents"),
      ("Lemon","Holliday - Immediate",2,"TBD","+1 to confirm"),
      ("Taylor","Holliday - Immediate",2,"TBD","+1 to confirm"),
      ("Ken & Ashley","Ken & Ashley",2,"No",""),
      ("Cody","Ken & Ashley",2,"TBD","+1 to confirm"),
      ("Kendal","Ken & Ashley",2,"TBD","+1 to confirm"),
      ("Jordan & Brad","Ken & Ashley",2,"No",""),
      ("Steve & Kathy","Steve & Kathy",2,"No",""),
      ("Nathan","Steve & Kathy",2,"TBD","+1 to confirm"),
      ("Alex","Steve & Kathy",1,"No","")]
T2 = [("Tom Rhodes","",2,"No",""), ("Kourtney & Andrew Wexler","",1,"No",""),
      ("Lane & Traci Jones","",2,"No",""), ("Bill & Terry Willits","",2,"No",""),
      ("Steve & Elaine Franklin","",2,"No",""), ("Andy & Sandra Stanley","",2,"No",""),
      ("Sherri & Wes Shropshire","",2,"No","")]
T3 = [("Jim & Shelley Sexton","",2,"No","Just thoughts"),
      ("Steve & Jana Harmon","",2,"No","Just thoughts"),
      ("Jeff Johnson & wife","",2,"No","Spouse name TBD"),
      ("Dave Tardella & wife","",2,"No","Spouse name TBD"),
      ("Brad Bretz & wife","",2,"No","Spouse name TBD"),
      ("Perko & Kristen","",2,"No",""),
      ("Network pastors (TBD)","",None,"No","Placeholder - decide who to add")]

EMMA_FAM = ["Will","Eve","Mia","Zoe","MaeMae","Mae Belle","Steve","Josh","Naomi","Elise",
            "Caleb","Michael","Clyde","Carly","Mathias","Rose","Wes","Katie","Brice","Marcus",
            "Ben","Catherine","Baby","Anna","Carrie","Dutch","Katherine","Jackson","Rebecca Mae",
            "Bear","Maggie","Kirkley","Margaret","Taylor","Crew","Wells","Elizabeth","Levi",
            "Ansen","Amy","Paul","Matthew","Lauren","Bellamy","Navy","Romie","Ashley","Bart",
            "Joan","Anna","Nick","Brayden","Thatcher","Ali","Kenny","Andre","Baby 2","JB",
            "Anabelle","Eli","Evie"]
EMMA_FAM_NOTES = {"Baby": "Name to confirm", "Ansen": "New baby - name to confirm",
                  "Brayden": "Name to confirm", "Baby 2": "Name to confirm"}
EMMA_FRIENDS = ["Denise Moss","Mike","Becky Miller","Matt","Donna Smith","Jackie","Lauren Hight",
                "Krin Baer","Rachel","David Wilhite","Lauren W","Clark","Martha","Dick",
                "Dick's wife","Chris","Chris' wife","Nikki Kintner","Pam Flowe","Anne Cundiff",
                "Tim Cundiff","Jim Glover","Anne Herbert","Jack Herbert"]
EMMA_FRIENDS_NOTES = {"Dick's wife": "Name to confirm", "Chris' wife": "Name to confirm"}

NEW = [("Brady Laubach","Laubach","No",""), ("Shannon Laubach","Laubach","No",""),
       ("Garrett Thurmond","Thurmond","No",""), ("Abby Thurmond","Thurmond","No",""),
       ("Collier Jackson","Jackson","No",""), ("Claire Jackson","Jackson","No",""),
       ("Bronson Lane","Lane","No",""), ("Rachel Lane","Lane","No",""),
       ("Jay Joyner","Joyner (Jay)","No",""),
       ("Jay Joyner's wife","Joyner (Jay)","No","Name to confirm"),
       ("Tyler Joyner","Joyner (Tyler)","No",""), ("Natalie Joyner","Joyner (Tyler)","No",""),
       ("Josh Havrilla","Havrilla","No",""),
       ("Josh Havrilla's +1","Havrilla","TBD","+1 to confirm"),
       ("Thomas Oden","Oden","No",""), ("Kaitlyn Oden","Oden","No",""),
       ("Zander","Zander","No","Last name to confirm"),
       ("Zander's wife","Zander","No","Name to confirm"),
       ("Avery","Avery","No","Last name to confirm"),
       ("Avery's wife","Avery","No","Name to confirm"),
       ("Walker Bradley","Bradley","No",""), ("Jordan Bradley","Bradley","No",""),
       ("Jon Barnett","Barnett","No",""),
       ("Abby / Maria","Barnett","No","Abby or Maria? Confirm name, and that she is Jon's guest"),
       ("Jack Baker","Baker","No",""),
       ("Andrew","Andrew","No","Last name to confirm - not the Andrew Wexler on row 12"),
       ("Aaron","Aaron","No","Last name to confirm"),
       ("Nick Stonaker","Stonaker","No",""),
       ("Nick Stonaker's +1","Stonaker","TBD","+1 to confirm"),
       ("Yoshi Nomura","Nomura","No",""), ("Yuka Nomura","Nomura","No",""),
       ("Rocky","Rocky","No","Last name to confirm"),
       ("Rocky's wife","Rocky","No","Name to confirm"),
       ("Shun","Shun & Momoka","No","Last name to confirm"),
       ("Momoka","Shun & Momoka","No","Last name to confirm"),
       ("David Wadijija","Wadijija","No","Spelling to confirm"),
       ("David Wadijija's +1","Wadijija","TBD","+1 to confirm"),
       ("[Placeholder - name TBD]","Placeholder","No","Listed as 'ZZZZZZZ' - decide who to add"),
       ("[Placeholder]'s +1","Placeholder","TBD","+1 to confirm"),
       ("Michael Rumende","Rumende","No",""),
       ("Chitra","Rumende","No","Last name to confirm"),
       ("PD","PD","No","Full name to confirm"),
       ("PD's wife","PD","No","Name to confirm")]


# Emma's friends - added after the Holliday-side friends
EMMA_FRIENDS_NEW = [
    ("Nathaniel", "Nathaniel & Rachel", "No", ""),
    ("Rachel", "Nathaniel & Rachel", "No", "Last name to confirm - not Rachel Lane or Rachel Guy"),
    ("Viet", "Viet & Debbie", "No", "Last name to confirm"),
    ("Debbie", "Viet & Debbie", "No", "Last name to confirm"),
    ("Noah", "Noah & Emily", "No", "Last name to confirm"),
    ("Emily", "Noah & Emily", "No", "Last name to confirm"),
    ("Michael", "Michael & Alisa", "No", "Last name to confirm - not Michael Rumende or Michael on row 36"),
    ("Alisa", "Michael & Alisa", "No", "Last name to confirm"),
    ("Grant", "Grant & Samantha", "No", "Last name to confirm"),
    ("Samantha", "Grant & Samantha", "No", "Last name to confirm"),
    ("Garrett", "Garrett & Moriah", "No", "Last name to confirm - not Garrett Thurmond"),
    ("Moriah", "Garrett & Moriah", "No", "Last name to confirm"),
    ("Madeleine Harris", "Harris", "No", ""),
    ("Bee Icayan", "Icayan", "No", ""),
    ("Anna Ruth Flagg", "Flagg", "No", ""),
    ("Rhiannon Beard", "Beard", "No", ""),
    ("Rhiannon Beard's +1", "Beard", "TBD", "+1 to confirm"),
    ("Megan Rogers", "Rogers", "No", ""),
    ("Norma Street", "Street", "No", ""),
    ("Jay Street", "Street", "No", ""),
    ("Jackie Santander", "Santander", "No", "Possible overlap with Jackie on row 91 - check"),
    ("Chris Santander", "Santander", "No", "Possible overlap with Chris / Chris' wife on rows 101-102 - check"),
    ("Sheona", "Sheona & Issac", "No", "Last name to confirm"),
    ("Issac", "Sheona & Issac", "No", "Spelling to confirm - Issac or Isaac?"),
    ("Rachel Guy", "Guy", "No", ""),
]

guests = []
for n, h, s, p, note in T1:
    guests.append([n, "Holliday", "Family (Tier 1)", h, s, p, note])
for n, h, s, p, note in T2:
    guests.append([n, "Holliday", "Framily (Tier 2)", h, s, p, note])
for n, h, s, p, note in T3:
    guests.append([n, "Holliday", "Influences (Tier 3)", h, s, p, note])
for n in EMMA_FAM:
    guests.append([n, "Emma", "Family", "Emma - Family", 1, "No", EMMA_FAM_NOTES.get(n, "")])
for n in EMMA_FRIENDS:
    guests.append([n, "Emma", "Parents' Friends", "Emma - Parents' Friends", 1, "No",
                   EMMA_FRIENDS_NOTES.get(n, "")])
guests.append(["Emma", "Couple", "Bride & Groom", "Bride & Groom", 1, "No",
               "Bride - delete this row if the counts should mean invited guests only"])
guests.append(["Matt", "Couple", "Bride & Groom", "Bride & Groom", 1, "No",
               "Groom - delete this row if the counts should mean invited guests only"])
for n, h, p, note in NEW:
    guests.append([n, "Holliday", "Friends (Tier 4)", h, 1, p, note])
for n, h, p, note in EMMA_FRIENDS_NEW:
    guests.append([n, "Emma", "Friends", h, 1, p, note])

title_block(gl, "GUEST LIST", "Your full list, organized by side and tier. Adjust party sizes as +1s are confirmed.", len(GL_HEAD))
header_row(gl, 3, GL_HEAD)

group_seen = []
for i, (name, side, cat, house, size, plus, note) in enumerate(guests):
    r = 4 + i
    key = house or cat
    if key not in group_seen:
        group_seen.append(key)
    gl.cell(row=r, column=1, value=i + 1)
    gl.cell(row=r, column=2, value=name)
    gl.cell(row=r, column=3, value=side)
    gl.cell(row=r, column=4, value=cat)
    gl.cell(row=r, column=5, value=house)
    gl.cell(row=r, column=6, value=size)
    gl.cell(row=r, column=7, value=plus)
    gl.cell(row=r, column=11, value="No")
    gl.cell(row=r, column=13, value="Pending")
    gl.cell(row=r, column=18, value=note)
    band(gl, r, len(GL_HEAD), group_seen.index(key) % 2 == 0)
    gl.cell(row=r, column=2).font = Font(color=TEXT, size=10, bold=True)
    gl.cell(row=r, column=18).font = Font(color=MID, size=9, italic=True)

GL_LAST = 3 + len(guests)
tot = GL_LAST + 1
gl.cell(row=tot, column=2, value="TOTALS").font = Font(bold=True, color="FFFFFF")
gl.cell(row=tot, column=6, value=f"=SUM(F4:F{GL_LAST})").font = Font(bold=True, color="FFFFFF")
gl.cell(row=tot, column=14, value=f"=SUM(N4:N{GL_LAST})").font = Font(bold=True, color="FFFFFF")
for c in range(1, len(GL_HEAD) + 1):
    gl.cell(row=tot, column=c).fill = PatternFill("solid", fgColor=MID)
    gl.cell(row=tot, column=c).border = BORDER

widths(gl, dict(zip("ABCDEFGHIJKLMNOPQR",
       [6, 27, 11, 18, 20, 10, 8, 20, 14, 22, 12, 11, 11, 11, 13, 13, 12, 50])))
gl.freeze_panes = "C4"
gl.auto_filter.ref = f"A3:{get_column_letter(len(GL_HEAD))}{GL_LAST}"

yn  = DataValidation(type="list", formula1='"Yes,No,TBD"', allow_blank=True)
rsv = DataValidation(type="list", formula1='"Pending,Yes,No"', allow_blank=True)
gl.add_data_validation(yn); gl.add_data_validation(rsv)
yn.add(f"G4:G{GL_LAST}"); yn.add(f"K4:L{GL_LAST}"); yn.add(f"P4:Q{GL_LAST}")
rsv.add(f"M4:M{GL_LAST}")
gl.conditional_formatting.add(f"M4:M{GL_LAST}", CellIsRule(
    operator="equal", formula=['"Yes"'], fill=PatternFill("solid", bgColor=BAND),
    font=Font(color=DEEP, bold=True)))
gl.conditional_formatting.add(f"M4:M{GL_LAST}", CellIsRule(
    operator="equal", formula=['"No"'], font=Font(color="9A3412")))

# ================================================================ OVERVIEW
ov = wb.create_sheet("Overview", 0)
title_block(ov, "WEDDING PLANNING WORKBOOK", "Overview - everything updates automatically from the other tabs.", 4)

def section(ws, row, label, span=4):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row=row, column=1, value=label)
    c.font = Font(bold=True, size=11, color="FFFFFF")
    c.alignment = Alignment(vertical="center", indent=1)
    for col in range(1, span + 1):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=MID)
        ws.cell(row=row, column=col).border = BORDER
    ws.row_dimensions[row].height = 22

def kv(ws, row, label, value, fmt=None, shade=True):
    a = ws.cell(row=row, column=1, value=label)
    b = ws.cell(row=row, column=2, value=value)
    a.font = Font(color=TEXT, size=10)
    b.font = Font(bold=True, color=DEEP, size=10)
    if fmt:
        b.number_format = fmt
    for col in (1, 2):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=BAND if shade else BAND2)
        ws.cell(row=row, column=col).border = BORDER

G = f"'Guest List'!"
section(ov, 4, "EVENT DETAILS")
kv(ov, 5, "Wedding Date", date(2026, 5, 22), "mmmm d, yyyy")
kv(ov, 6, "Venue", "", None, False)
kv(ov, 7, "Total Budget Goal", 25000, MONEY)
kv(ov, 8, "Days Until Wedding", "=B5-TODAY()", "0", False)
ov.cell(row=8, column=3, value="Negative = the date has passed; update B5 if the year is wrong").font = Font(color=MID, size=9, italic=True)

section(ov, 10, "GUEST SUMMARY")
kv(ov, 11, "Total Invited Parties", f"=COUNTA({G}B4:B{GL_LAST})")
kv(ov, 12, "Total Invited Guests", f"=SUM({G}F4:F{GL_LAST})", None, False)
kv(ov, 13, "RSVP - Yes (parties)", f'=COUNTIF({G}M4:M{GL_LAST},"Yes")')
kv(ov, 14, "RSVP - No (parties)", f'=COUNTIF({G}M4:M{GL_LAST},"No")', None, False)
kv(ov, 15, "RSVP - Pending (parties)", f'=COUNTIF({G}M4:M{GL_LAST},"Pending")')
kv(ov, 16, "Total Attending (guests)", f"=SUM({G}N4:N{GL_LAST})", None, False)

section(ov, 18, "GUESTS BY GROUP")
for i, h in enumerate(["Side", "Category", "Parties", "Guests"], start=1):
    c = ov.cell(row=19, column=i, value=h)
    c.font = Font(bold=True, color="FFFFFF", size=10)
    c.fill = PatternFill("solid", fgColor=DEEP)
    c.border = BORDER
GROUPS = [("Holliday", "Family (Tier 1)"), ("Holliday", "Framily (Tier 2)"),
          ("Holliday", "Influences (Tier 3)"), ("Holliday", "Friends (Tier 4)"),
          ("Emma", "Family"), ("Emma", "Parents' Friends"), ("Emma", "Friends"),
          ("Couple", "Bride & Groom")]
for i, (side, cat) in enumerate(GROUPS):
    r = 20 + i
    ov.cell(row=r, column=1, value=side)
    ov.cell(row=r, column=2, value=cat)
    ov.cell(row=r, column=3, value=f'=COUNTIFS({G}C4:C{GL_LAST},A{r},{G}D4:D{GL_LAST},B{r})')
    ov.cell(row=r, column=4, value=f'=SUMIFS({G}F4:F{GL_LAST},{G}C4:C{GL_LAST},A{r},{G}D4:D{GL_LAST},B{r})')
    band(ov, r, 4, i % 2 == 0)
gr = 20 + len(GROUPS)
ov.cell(row=gr, column=1, value="TOTAL")
ov.cell(row=gr, column=3, value=f"=SUM(C20:C{gr-1})")
ov.cell(row=gr, column=4, value=f"=SUM(D20:D{gr-1})")
for c in range(1, 5):
    ov.cell(row=gr, column=c).fill = PatternFill("solid", fgColor=MID)
    ov.cell(row=gr, column=c).font = Font(bold=True, color="FFFFFF", size=10)
    ov.cell(row=gr, column=c).border = BORDER

B = "Budget!"
section(ov, gr + 2, "BUDGET SNAPSHOT")
kv(ov, gr + 3, "Estimated Total", f"=SUM({B}D4:D31)", MONEY)
kv(ov, gr + 4, "Actual Total", f"=SUM({B}E4:E31)", MONEY, False)
kv(ov, gr + 5, "Paid to Date", f"=SUM({B}F4:F31)", MONEY)
kv(ov, gr + 6, "Balance Remaining", f"=SUM({B}G4:G31)", MONEY, False)
kv(ov, gr + 7, "Budget Goal vs Estimate", f"=B7-SUM({B}D4:D31)", MONEY)

C = "'Planning Checklist'!"
section(ov, gr + 9, "CHECKLIST PROGRESS")
kv(ov, gr + 10, "Tasks - Done", f'=COUNTIF({C}F4:F44,"Done")')
kv(ov, gr + 11, "Tasks - In Progress", f'=COUNTIF({C}F4:F44,"In Progress")', None, False)
kv(ov, gr + 12, "Tasks - Not Started", f'=COUNTIF({C}F4:F44,"Not Started")')
kv(ov, gr + 13, "% Complete", f'=IFERROR(COUNTIF({C}F4:F44,"Done")/COUNTA({C}A4:A44),0)', "0.0%", False)
widths(ov, {"A": 26, "B": 22, "C": 12, "D": 12})

# ================================================================ BUDGET
bd = wb.create_sheet("Budget")
BD_HEAD = ["Category", "Item / Description", "Vendor", "Estimated ($)", "Actual ($)",
           "Paid ($)", "Balance ($)", "Due Date", "Paid?", "Notes"]
ITEMS = [("Venue & Rentals", "Ceremony & reception site"), ("Venue & Rentals", "Tables, chairs & linens"),
         ("Venue & Rentals", "Tent / lighting / restrooms"), ("Catering & Bar", "Food (per head)"),
         ("Catering & Bar", "Bar & beverages"), ("Catering & Bar", "Service staff & gratuity"),
         ("Catering & Bar", "Cake / desserts"), ("Attire & Beauty", "Wedding dress & alterations"),
         ("Attire & Beauty", "Suit / tux"), ("Attire & Beauty", "Hair & makeup"),
         ("Attire & Beauty", "Accessories & shoes"), ("Photo & Video", "Photographer"),
         ("Photo & Video", "Videographer"), ("Flowers & Decor", "Bouquets & boutonnieres"),
         ("Flowers & Decor", "Ceremony & reception florals"), ("Flowers & Decor", "Centerpieces & signage"),
         ("Entertainment", "Ceremony musicians"), ("Entertainment", "DJ / band"),
         ("Stationery", "Save-the-dates & invitations"), ("Stationery", "Programs, menus & thank-yous"),
         ("Rings & Gifts", "Wedding bands"), ("Rings & Gifts", "Party gifts & favors"),
         ("Officiant & Fees", "Officiant fee"), ("Officiant & Fees", "Marriage license"),
         ("Transport & Stay", "Couple & party transport"), ("Transport & Stay", "Hotel block / night-of room"),
         ("Planning", "Coordinator / planner"), ("Other", "Contingency (~10%)")]
title_block(bd, "BUDGET", "Enter numbers in the Estimated / Actual / Paid columns - balances and totals calculate themselves.", len(BD_HEAD))
header_row(bd, 3, BD_HEAD)
cats = []
for i, (cat, item) in enumerate(ITEMS):
    r = 4 + i
    if cat not in cats:
        cats.append(cat)
    bd.cell(row=r, column=1, value=cat)
    bd.cell(row=r, column=2, value=item)
    bd.cell(row=r, column=7, value=f"=IF(N(E{r})+N(F{r})=0,\"\",N(E{r})-N(F{r}))")
    band(bd, r, len(BD_HEAD), cats.index(cat) % 2 == 0)
    for c in (4, 5, 6, 7):
        bd.cell(row=r, column=c).number_format = MONEY
BD_LAST = 3 + len(ITEMS)
r = BD_LAST + 1
bd.cell(row=r, column=1, value="TOTAL")
for c, col in ((4, "D"), (5, "E"), (6, "F"), (7, "G")):
    bd.cell(row=r, column=c, value=f"=SUM({col}4:{col}{BD_LAST})").number_format = MONEY
for c in range(1, len(BD_HEAD) + 1):
    bd.cell(row=r, column=c).fill = PatternFill("solid", fgColor=MID)
    bd.cell(row=r, column=c).font = Font(bold=True, color="FFFFFF", size=10)
    bd.cell(row=r, column=c).border = BORDER
widths(bd, dict(zip("ABCDEFGHIJ", [18, 30, 20, 14, 13, 12, 13, 12, 9, 34])))
bd.freeze_panes = "C4"
paid = DataValidation(type="list", formula1='"Yes,No,Partial"', allow_blank=True)
bd.add_data_validation(paid); paid.add(f"I4:I{BD_LAST}")

# ================================================================ VENDORS
vn = wb.create_sheet("Vendors & Contacts")
VN_HEAD = ["Vendor Type", "Company / Name", "Contact Person", "Phone", "Email", "Website",
           "Quoted ($)", "Deposit ($)", "Balance ($)", "Contract Signed", "Notes"]
TYPES = ["Venue", "Caterer", "Bar Service", "Photographer", "Videographer", "Florist",
         "DJ / Band", "Ceremony Musicians", "Officiant", "Cake / Baker", "Hair & Makeup",
         "Planner / Coordinator", "Rentals", "Transportation", "Stationery", "Hotel Block"]
title_block(vn, "VENDORS & CONTACTS", "Keep every quote, contact and contract in one place.", len(VN_HEAD))
header_row(vn, 3, VN_HEAD)
for i, t in enumerate(TYPES):
    r = 4 + i
    vn.cell(row=r, column=1, value=t)
    vn.cell(row=r, column=9, value=f"=IF(N(G{r})+N(H{r})=0,\"\",N(G{r})-N(H{r}))")
    band(vn, r, len(VN_HEAD), i % 2 == 0)
    for c in (7, 8, 9):
        vn.cell(row=r, column=c).number_format = MONEY
VN_LAST = 3 + len(TYPES)
widths(vn, dict(zip("ABCDEFGHIJK", [22, 26, 20, 16, 26, 24, 12, 12, 12, 14, 34])))
vn.freeze_panes = "B4"
sig = DataValidation(type="list", formula1='"Yes,No,Pending"', allow_blank=True)
vn.add_data_validation(sig); sig.add(f"J4:J{VN_LAST}")

# ================================================================ CHECKLIST
ck = wb.create_sheet("Planning Checklist")
CK_HEAD = ["Task", "Category", "Timeframe", "Due Date", "Owner", "Status", "Notes"]
TASKS = [("Set a wedding budget","Budget","12+ Months Out"),("Draft the guest list","Guests","12+ Months Out"),
    ("Pick a wedding date","Planning","12+ Months Out"),("Choose & book the venue","Venue","12+ Months Out"),
    ("Hire a planner (optional)","Planning","12+ Months Out"),("Book photographer & videographer","Vendors","9-11 Months Out"),
    ("Book caterer","Vendors","9-11 Months Out"),("Book band / DJ","Vendors","9-11 Months Out"),
    ("Order wedding dress","Attire","9-11 Months Out"),("Reserve hotel room block","Lodging","9-11 Months Out"),
    ("Send save-the-dates","Stationery","6-8 Months Out"),("Book florist","Vendors","6-8 Months Out"),
    ("Book officiant","Vendors","6-8 Months Out"),("Plan ceremony details","Planning","6-8 Months Out"),
    ("Arrange transportation","Vendors","6-8 Months Out"),("Register for gifts","Planning","6-8 Months Out"),
    ("Order invitations","Stationery","4-5 Months Out"),("Choose & order cake","Catering","4-5 Months Out"),
    ("Finalize menu & tasting","Catering","4-5 Months Out"),("Buy / rent suits","Attire","4-5 Months Out"),
    ("Schedule hair & makeup trial","Beauty","4-5 Months Out"),("Mail invitations","Stationery","2-3 Months Out"),
    ("Finalize readings & vows","Planning","2-3 Months Out"),("Purchase wedding bands","Rings","2-3 Months Out"),
    ("Apply for marriage license","Legal","2-3 Months Out"),("Confirm details with vendors","Vendors","2-3 Months Out"),
    ("Track RSVPs & follow up","Guests","1 Month Out"),("Create seating chart","Guests","1 Month Out"),
    ("Final dress fitting","Attire","1 Month Out"),("Confirm headcount with caterer","Catering","1 Month Out"),
    ("Confirm timeline with vendors & party","Planning","1 Week Out"),("Prepare final payments & tips","Budget","1 Week Out"),
    ("Delegate day-of tasks","Planning","1 Week Out"),("Pack for honeymoon","Travel","1 Week Out"),
    ("Eat breakfast & stay hydrated","Self-care","Day Of"),("Hair & makeup","Beauty","Day Of"),
    ("Exchange / give party gifts","Gifts","Day Of"),("Enjoy every moment!","Self-care","Day Of"),
    ("Send thank-you notes","Guests","After"),("Return rentals","Vendors","After"),
    ("Review / tip vendors","Vendors","After")]
title_block(ck, "PLANNING CHECKLIST", "A full timeline from 'just engaged' to 'I do' - update Status as you go.", len(CK_HEAD))
header_row(ck, 3, CK_HEAD)
frames = []
for i, (task, cat, frame) in enumerate(TASKS):
    r = 4 + i
    if frame not in frames:
        frames.append(frame)
    ck.cell(row=r, column=1, value=task)
    ck.cell(row=r, column=2, value=cat)
    ck.cell(row=r, column=3, value=frame)
    ck.cell(row=r, column=6, value="Not Started")
    band(ck, r, len(CK_HEAD), frames.index(frame) % 2 == 0)
CK_LAST = 3 + len(TASKS)
widths(ck, dict(zip("ABCDEFG", [38, 14, 18, 13, 14, 15, 36])))
ck.freeze_panes = "B4"
st = DataValidation(type="list", formula1='"Not Started,In Progress,Done"', allow_blank=True)
ck.add_data_validation(st); st.add(f"F4:F{CK_LAST}")
ck.conditional_formatting.add(f"F4:F{CK_LAST}", CellIsRule(
    operator="equal", formula=['"Done"'], fill=PatternFill("solid", bgColor=BAND),
    font=Font(color=DEEP, bold=True)))

# ================================================================ SEATING
sc = wb.create_sheet("Seating Chart")
SC_HEAD = ["Table #", "Guest / Party Name", "# Seats", "Side", "Meal", "Notes"]
title_block(sc, "SEATING CHART", "Assign each party to a table, then sort by Table # to see your layout.", len(SC_HEAD))
header_row(sc, 3, SC_HEAD)
SC_LAST = 3 + 60
for i in range(60):
    band(sc, 4 + i, len(SC_HEAD), (i // 2) % 2 == 0)
r = SC_LAST + 1
sc.cell(row=r, column=2, value="Seats Assigned")
sc.cell(row=r, column=3, value=f"=SUM(C4:C{SC_LAST})")
for c in range(1, len(SC_HEAD) + 1):
    sc.cell(row=r, column=c).fill = PatternFill("solid", fgColor=MID)
    sc.cell(row=r, column=c).font = Font(bold=True, color="FFFFFF", size=10)
    sc.cell(row=r, column=c).border = BORDER
widths(sc, dict(zip("ABCDEF", [10, 30, 10, 12, 16, 40])))
sc.freeze_panes = "B4"

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = DEEP
    ws.sheet_view.showGridLines = False

wb.save(OUT)
print("saved", OUT)
print("guest rows:", len(guests), "| last guest row:", GL_LAST)
